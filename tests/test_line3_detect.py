"""LINE3 階段一：批次一讀檔／表頭判定與批次二 Operating Modes 解析。"""
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.cell_range import CellRange

from alarm_ingest.detect import (
    _cell_to_str, _find_precise_header_row, _normalize_header_cell, detect_columns,
    list_sheets, read_grid, read_tabular, grid_to_rows,
    _find_operating_modes_range, _extract_operating_modes_sub_fields,
    extract_operating_modes, _parse_operating_modes_cell, OperatingModesCellState,
)

HEADERS = [
    'Alarm Type', 'Alarm Number', 'HMI Message', 'Alarm Description',
    'Solution', 'Alarm Simulation', 'Alarm Effects', 'Alarm Reset',
    'Validation Test', 'Operating Modes',
]
ROOT = Path(__file__).resolve().parents[1]
FILL = ROOT / 'tools/variant/fixtures/FILL203_batch_report_alarm_list_260529_ENG_revise.xlsx'


def _line3_fixture_path(filename):
    """兩批共用原檔定位；指定目錄卻缺檔時必須失敗，不能 skip。"""
    names = [filename, filename.replace('.xlsx', ' (Alarms List).xlsx')]
    fixture_dir = os.environ.get('LINE3_FIXTURE_DIR')
    if fixture_dir:
        candidates = [Path(fixture_dir) / name for name in names]
    else:
        candidates = [path for name in names for path in ROOT.rglob(name)]
        candidates.extend(ROOT.parent / 'LINE 3' / name for name in names)
    for path in candidates:
        if path.is_file():
            return path
    if fixture_dir:
        pytest.fail(f'Missing required fixture: {filename} in {fixture_dir}')
    pytest.skip('LINE3 原始 Excel 未提供；請設定 LINE3_FIXTURE_DIR')


def data(number):
    return ['A', number, 'message', 'description', 'solution']


def test_normalization():
    assert _normalize_header_cell('  ＡＬＡＲＭ\n Ｎｕｍｂｅｒ　') == 'alarm number'
    assert _normalize_header_cell(None) == ''


def test_grid_merges_and_sheet_selection(tmp_path):
    path = tmp_path / 'merged.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'alarms'
    ws.append(HEADERS)
    ws.merge_cells('J1:L1')
    ws.merge_cells('A1:A2')
    wb.create_sheet('empty')
    wb.save(path)
    wb.close()
    assert all(len(item) == 3 and item[2] == [] for item in read_grid(path))
    direct = openpyxl.load_workbook(path, data_only=True)
    try:
        actual = read_grid(path, with_merges=True)
        assert [(n, set(map(str, m))) for n, _, m in actual] == [
            (ws.title, set(map(str, ws.merged_cells.ranges))) for ws in direct
        ]
    finally:
        direct.close()
    assert len(read_grid(path, sheet='alarms', with_merges=True)) == 1
    with pytest.raises(ValueError, match='找不到分頁'):
        read_grid(path, sheet='missing')


def test_csv_and_existing_callers(tmp_path):
    path = tmp_path / 'old.csv'
    path.write_text('Description,Cause,Action\n' + ''.join(
        f'{100+i} message,cause,action\n' for i in range(5)), encoding='utf-8-sig')
    assert read_grid(path, with_merges=True)[0][2] == []
    assert read_grid(path)[0][2] == []
    assert list_sheets(path) == [('csv', True)]
    assert len(read_tabular(path)) == 5


@pytest.mark.parametrize('row_number', [58, 65, 75])
def test_deep_header_synthetic(row_number):
    grid = [['title']] * (row_number - 1) + [HEADERS, [], data(2945)]
    assert _find_precise_header_row(grid) == row_number - 1
    assert detect_columns(grid) == (1, 3, 4, row_number + 1)


@pytest.mark.parametrize('numbers,expected', [
    ([1, 2, 3, 4, 'bad'], 0),
    ([1, 2, 3, 'bad', 'bad'], None),
    ([0.0], 0), ([1.5], None), ([], None), (['123x'], None),
])
def test_condition_b_threshold(numbers, expected):
    grid = [HEADERS, []] + [data(n) for n in numbers] + [[], [None, '  ']]
    assert _find_precise_header_row(grid) == expected


def test_condition_a_exact_and_minimum():
    assert _find_precise_header_row([HEADERS[:6], [], data(1)]) == 0
    assert _find_precise_header_row([HEADERS[:5], [], data(1)]) is None
    assert _find_precise_header_row([[h + ' explanation' for h in HEADERS], [], data(1)]) is None
    missing_number = [h for h in HEADERS if h != 'Alarm Number']
    assert _find_precise_header_row([missing_number, [], data(1)]) is None


def test_legend_rejected_by_b_then_later_header_found():
    grid = [HEADERS, []] + [['legend', 'not a number']] * 8
    grid += [HEADERS, [], data(123)]
    assert _find_precise_header_row(grid) == 10


def test_condition_c_later_header_wins_tie():
    # Both candidates pass B: the earlier one has 5/6 numeric rows.
    grid = [HEADERS, [], HEADERS, []] + [data(n) for n in range(5)]
    assert _find_precise_header_row(grid) == 2


def test_condition_c_more_hits_win_before_later_row():
    grid = [HEADERS, [], HEADERS[:6], []] + [data(n) for n in range(5)]
    assert _find_precise_header_row(grid) == 0


def test_fill203_real_fixture_unchanged():
    expected = {'alarm count': (1, 2, 3, 1), 'alarm list': (0, 1, 2, 2),
                'Problem': None, '工作表1': None}
    grids = read_grid(FILL)
    assert {name: detect_columns(grid) for name, grid, _ in grids} == expected
    assert all(_find_precise_header_row(grid) is None for _, grid, _ in grids)
    assert list_sheets(FILL) == [(name, result is not None) for name, result in expected.items()]
    assert read_tabular(FILL, sheet='alarm list')


@pytest.mark.parametrize('filename,row_number', [
    ('22062QE6SP63_LF.xlsx', 58), ('22062QE5SP63_02.xlsx', 65),
    ('22062QE6SP63_DE.xlsx', 75), ('22062QE6SP63_RI.xlsx', 75),
    ('22062QE6SP63_DB_01.xlsx', 75),
])
def test_line3_real_fixture(filename, row_number):
    path = _line3_fixture_path(filename)
    grids = read_grid(path, with_merges=True)
    assert any(_find_precise_header_row(grid) == row_number - 1 for _, grid, _ in grids)
    direct = openpyxl.load_workbook(path, data_only=True)
    try:
        for name, grid, merges in grids:
            assert set(map(str, merges)) == set(map(str, direct[name].merged_cells.ranges))
            if _find_precise_header_row(grid) == row_number - 1:
                assert detect_columns(grid)[3] == row_number + 1
    finally:
        direct.close()


@pytest.mark.parametrize('filename,names,count', [
    ('22062QE6SP63_LF.xlsx', ['PROD', 'MM'], 4),
    ('22062QE5SP63_02.xlsx', ['PM', 'ES', 'LT', 'VHP', 'EA', 'CL', 'MM'], 92),
    ('22062QE6SP63_DE.xlsx', ['PREP', 'PROD', 'PE', 'MM', 'CS', 'VHP'], 20),
    ('22062QE6SP63_RI.xlsx', ['PREP', 'PROD', 'PE', 'MM', 'CS', 'VHP'], 95),
    ('22062QE6SP63_DB_01.xlsx', ['PREP', 'PROD', 'PE', 'MM', 'CS', 'VHP'], 3),
])
def test_operating_modes_real_fixture(filename, names, count):
    path = _line3_fixture_path(filename)
    converted = []
    for sheet, grid, merges in read_grid(path, with_merges=True):
        header_idx = _find_precise_header_row(grid)
        if header_idx is None:
            continue
        col_range = _find_operating_modes_range(header_idx, grid[header_idx], merges)
        assert col_range == (9, 9 + len(names) - 1)
        sub_fields = _extract_operating_modes_sub_fields(grid, header_idx, col_range)
        assert sub_fields == names
        mapping = detect_columns(grid)
        rows = grid_to_rows(grid, mapping, filename + '#' + sheet)
        indices = [i for i in range(mapping[3], len(grid))
                   if _cell_to_str(grid[i][1]).strip()]
        assert len(rows) == len(indices)
        for row, idx in zip(rows, indices):
            assert row['code'] == _cell_to_str(grid[idx][1]).strip()
            modes, warnings = extract_operating_modes(grid, header_idx, col_range, sub_fields, idx)
            assert warnings == []
            assert modes == [name for col, name in enumerate(names, 9)
                             if str(grid[idx][col]).strip().lower() == 'x']
            row['operating_modes'] = modes
        converted.extend(rows)
    assert len(converted) == count


def test_operating_modes_workbook_right_boundary_and_sparse_subheaders(tmp_path):
    path = tmp_path / 'trailing.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['title'])
    ws.append(['Code', ' ＯＰＥＲＡＴＩＮＧ\n Modes ', None, None, 'Other'])
    ws.merge_cells('B2:D2')
    ws.append([None, 'PROD', None, 'MM', 'Unrelated'])
    ws.append([123, '-', 'unnamed', 'x', 'outside range'])
    wb.save(path)
    wb.close()
    _, grid, merges = read_grid(path, with_merges=True)[0]
    col_range = _find_operating_modes_range(1, grid[1], merges)
    assert col_range == (1, 3)
    sub_fields = _extract_operating_modes_sub_fields(grid, 1, col_range)
    assert sub_fields == ['PROD', 'MM']
    assert extract_operating_modes(grid, 1, col_range, sub_fields, 3) == (['MM'], [])


@pytest.mark.parametrize('merges', [[], [CellRange('B1:B2')]])
def test_operating_modes_single_column(merges):
    grid = [['Code', 'Operating Modes', 'Other'], [None, 'MM', 'Unrelated'], [123, 'x', 'x']]
    col_range = _find_operating_modes_range(0, grid[0], merges)
    assert col_range == (1, 1)
    assert _extract_operating_modes_sub_fields(grid, 0, col_range) == ['MM']
    assert extract_operating_modes(grid, 0, col_range, ['MM'], 2) == (['MM'], [])


@pytest.mark.parametrize('values,expected', [
    (['x', 'X', ' Ｘ '], ['PROD', 'MM', 'CS']),
    (['-', 'x', 'n/a'], ['MM']),
    (['', '—', None], []), ([], []),
])
def test_operating_modes_representative_rows(values, expected):
    grid = [['Operating Modes'], ['PROD', 'MM', 'CS'], values]
    assert extract_operating_modes(grid, 0, (0, 2), grid[1], 2) == (expected, [])


def test_operating_modes_warning_details():
    grid = [['Operating Modes'], ['PROD', 'MM', 'CS', 'VHP'], ['x', ' yes ', 0, False]]
    assert extract_operating_modes(grid, 0, (0, 3), grid[1], 2) == (['PROD'], [
        {'field': 'MM', 'row': 2, 'column': 1, 'raw_value': ' yes '},
        {'field': 'CS', 'row': 2, 'column': 2, 'raw_value': '0'},
        {'field': 'VHP', 'row': 2, 'column': 3, 'raw_value': 'False'},
    ])


@pytest.mark.parametrize('value,state', [
    ('x', OperatingModesCellState.ENABLED),
    (' Ｘ\n', OperatingModesCellState.ENABLED),
    (None, OperatingModesCellState.DISABLED),
    (' \t', OperatingModesCellState.DISABLED),
    ('-', OperatingModesCellState.DISABLED),
    ('—', OperatingModesCellState.DISABLED),
    (' Ｎ／Ａ ', OperatingModesCellState.DISABLED),
    ('yes', OperatingModesCellState.UNKNOWN),
    (1.0, OperatingModesCellState.UNKNOWN),
    (False, OperatingModesCellState.UNKNOWN),
])
def test_operating_modes_three_states(value, state):
    assert _parse_operating_modes_cell(value) is state


def test_operating_modes_absence():
    grid = [['Operating Modes explanation'], ['MM'], ['-']]
    col_range = _find_operating_modes_range(0, grid[0], [])
    assert col_range is None
    assert extract_operating_modes(grid, 0, col_range, [], 2) == (None, [])
    assert extract_operating_modes(grid, 0, (0, 0), ['MM'], 2) == ([], [])


def test_operating_modes_invalid_mapping():
    assert _extract_operating_modes_sub_fields([['Operating Modes']], 0, (0, 1)) == []
    grid = [['Operating Modes'], ['MM'], ['x']]
    with pytest.raises(ValueError, match='子欄位數量'):
        extract_operating_modes(grid, 0, (0, 0), [], 2)
    for idx in (-1, 3):
        with pytest.raises(IndexError, match='資料列索引'):
            extract_operating_modes(grid, 0, (0, 0), ['MM'], idx)
