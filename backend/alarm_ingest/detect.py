"""原廠格式（路線 A）的欄位偵測。純函式，分三層：讀檔成 grid、偵測欄位、
依人工確認過的欄位對應轉成 rows——三層之間刻意不耦合，因為「偵測」與
「轉換」中間隔著一個必要的人工確認關卡（見批次匯入 UI 規劃第 3 節）：
inspect 端點只能做到偵測，不能在人還沒確認欄位對應之前就先把 code 切好，
否則「人工確認」這一步會變成看既成事實而非真正做決定。

跟 Variant/parse_alarms.py（現 tools/variant/parse_alarms.py）共用同一份
_detect_columns() 邏輯，而非各自維護一份複本——這份判斷比 normalize_variant
複雜得多（表頭比對＋fallback＋門檻＋兩項誤報防護），分岔的代價是同一份
原廠檔案在 CLI 與後台跑出不同的欄位對應，且契約測試很難完整覆蓋各種
grid 形狀去守住一致性，比 normalize_variant 的分岔風險高很多。CLI 的
read_tabular() 改為呼叫這裡的 read_grid()/detect_columns()/grid_to_rows()
組合，行為不變（CLI 自動採用 detect_columns() 的建議，等同人工確認
這一步在 CLI 上是自動通過——這符合 CLI 「你在跑、你知道自己在做什麼」
的定位）。
"""
from __future__ import annotations

import csv
import pathlib
import re
import sys
import unicodedata
from collections import Counter
from enum import Enum

from .quality import clean, split_code as _split_code

# 代碼在文字開頭：「31033 Operation active...」或「0024 - Forming Panel...」
CODE_PREFIX_RE = re.compile(r"^\s*(\d{3,6})\s*[-–—:：]?\s*(.*)$", re.S)


def split_code(text: str) -> tuple:
    """把「31033 Operation active weigh-in filling 1」拆成 (code, variant)。
    quality.split_code() 的特化版本，固定套用本模組的 CODE_PREFIX_RE
    （跟正則表達式耦合的部分留在這裡，quality.py 保持通用）。"""
    return _split_code(text, CODE_PREFIX_RE)


def _cell_to_str(v) -> str:
    """openpyxl 回傳的儲存格型別取決於儲存格格式，不保證是字串——
    code 這類欄位使用者可能存成數字（例如 24 而非 "24"），直接 str()
    對整數值沒問題，但對「整數值卻是 float 型別」的儲存格（公式結果、
    複製貼上運算、pandas 匯出等常見來源）會產生 "24.0" 這種帶小數尾碼
    的錯誤字串——這不是崩潰，是安靜產生錯的代碼，比 AttributeError
    更難發現。float 若無小數部分，先轉 int 再轉字串去掉尾碼；有小數
    部分（例如儲存格真的存了 31033.5）才保留原樣，讓後續驗證去處理
    這種本來就不像有效代碼格式的輸入。

    inspect 端點取樣本內容時也需要同樣的保護（避免 openpyxl 數字型別
    造成 JSON 序列化問題），跟 load_excel() 共用這支而非各自實作。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_grid(path: pathlib.Path, sheet: str = None, with_merges: bool = False) -> list:
    """讀 .xlsx/.xlsm/.csv 成 [(分頁名, grid, merges), ...]。純讀檔，不做欄位判斷、
    不做 code 切分。CSV 沒有分頁概念，統一包成一個 ("csv", grid, []) 維持
    介面一致。with_merges=True 時讀取 Excel 合併範圍，否則 merges 為空 list。

    sheet 指定時只回傳該分頁（找不到則報錯）——多分頁檔案預設會全部
    讀入，曾經因此把不相關的分頁一起解析進來而不自知（ACM002警報.xlsx
    有兩個分頁都含警報格式的資料，但資料庫裡只匯入過其中一個），所以
    需要能只鎖定單一分頁重現特定來源。"""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            grids = [("csv", [list(r) for r in csv.reader(f)], [])]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            grids = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)],
                      list(ws.merged_cells.ranges) if with_merges else [])
                     for ws in wb.worksheets]
        finally:
            wb.close()

    if sheet is None:
        return grids
    matched = [g for g in grids if g[0] == sheet]
    if not matched:
        available = ", ".join(repr(n) for n, _, _ in grids)
        raise ValueError(f"找不到分頁 {sheet!r}，此檔案的分頁：{available}")
    return matched


# LINE3 原廠模板的整格表頭；支援新片語時在此擴充並補上格式測試。
PRECISE_HEADER_PHRASES = frozenset({
    "alarm type", "alarm number", "hmi message", "alarm description",
    "solution", "alarm simulation", "alarm effects", "alarm reset",
    "validation test", "operating modes",
})


def _normalize_header_cell(value) -> str:
    """全形英文轉半形，將換行與連續空白合併，忽略大小寫。"""
    return " ".join(unicodedata.normalize("NFKC", _cell_to_str(value)).split()).casefold()


def _find_precise_header_row(grid: list):
    """回傳通過 A/B 的最佳表頭列索引；同分取較後列，找不到回 None。"""
    best = None
    for row_idx, row in enumerate(grid):
        cells = [_normalize_header_cell(value) for value in row]
        hits = sum(cell in PRECISE_HEADER_PHRASES for cell in cells)
        if hits < 6 or "alarm number" not in cells:
            continue
        number_col = cells.index("alarm number")
        nonempty = matches = 0
        for data_row in grid[row_idx + 2:]:
            if not any(_cell_to_str(value).strip() for value in data_row):
                continue
            nonempty += 1
            if number_col < len(data_row):
                number = _cell_to_str(data_row[number_col]).strip()
                matches += bool(re.fullmatch(r"\d+", number))
        if matches and matches * 5 >= nonempty * 4:
            candidate = (hits, row_idx)
            if best is None or candidate > best:
                best = candidate
    return best[1] if best is not None else None


def _find_operating_modes_range(header_row_idx: int, header_row: list, merges: list):
    """回傳 Operating Modes 的欄範圍（0-based，含右界）；缺席回 None。

    merges 必須來自 read_grid(..., with_merges=True)，不猜測右側欄位。
    """
    for col_idx, cell in enumerate(header_row):
        if _normalize_header_cell(cell) != "operating modes":
            continue
        for merged in merges:
            if (merged.min_row <= header_row_idx + 1 <= merged.max_row
                    and merged.min_col <= col_idx + 1 <= merged.max_col):
                return merged.min_col - 1, merged.max_col - 1
        return col_idx, col_idx
    return None


def _operating_modes_sub_columns(grid: list, header_row_idx: int, col_range: tuple) -> list:
    """保留非空子表頭的實際欄索引，避免中間空白造成後續欄位左移。"""
    sub_row = grid[header_row_idx + 1] if header_row_idx + 1 < len(grid) else []
    start_col, end_col = col_range
    return [(col, _cell_to_str(sub_row[col]).strip())
            for col in range(start_col, min(end_col + 1, len(sub_row)))
            if _cell_to_str(sub_row[col]).strip()]


def _extract_operating_modes_sub_fields(grid: list, header_row_idx: int,
                                       col_range: tuple) -> list:
    """依原始欄序回傳非空子欄位名稱，不排序、不補齊空白子表頭。"""
    return [name for _, name in _operating_modes_sub_columns(grid, header_row_idx, col_range)]


class OperatingModesCellState(Enum):
    ENABLED = "enabled"
    DISABLED = "disabled"
    UNKNOWN = "unknown"


def _parse_operating_modes_cell(value) -> OperatingModesCellState:
    """三態判斷；UNKNOWN 必須由呼叫端記錄警告，不可當成 DISABLED。"""
    normalized = _normalize_header_cell(value)
    if normalized == "x":
        return OperatingModesCellState.ENABLED
    if normalized in ("", "-", "—", "n/a"):
        return OperatingModesCellState.DISABLED
    return OperatingModesCellState.UNKNOWN


def extract_operating_modes(grid: list, header_row_idx: int, col_range,
                            sub_fields: list, data_row_idx: int) -> tuple:
    """解析一列，回傳 (啟用名稱清單, 警告清單)，不寫入或依 code 去重。

    col_range=None 表示檔案缺席此欄，回傳 (None, [])；有欄但全未啟用
    則回傳 ([], [])。sub_fields 由 _extract_operating_modes_sub_fields()
    取得，與非空子表頭位置一一對應。警告保留原值文字、欄名及 0-based
    row/column，呼叫端負責連同來源檔名/分頁呈現；有警告時清單只包含
    確認啟用的模式，不代表其他模式已確認停用。
    """
    if col_range is None:
        return None, []
    columns = _operating_modes_sub_columns(grid, header_row_idx, col_range)
    if len(columns) != len(sub_fields):
        raise ValueError("Operating Modes 子欄位數量與子表頭位置不符")
    if not 0 <= data_row_idx < len(grid):
        raise IndexError("Operating Modes 資料列索引超出範圍")
    row = grid[data_row_idx]
    modes, warnings = [], []
    for (col_idx, _), field_name in zip(columns, sub_fields):
        value = row[col_idx] if col_idx < len(row) else None
        state = _parse_operating_modes_cell(value)
        if state is OperatingModesCellState.ENABLED:
            modes.append(field_name)
        elif state is OperatingModesCellState.UNKNOWN:
            warnings.append({"field": field_name, "row": data_row_idx,
                             "column": col_idx, "raw_value": _cell_to_str(value)})
    return modes, warnings


def detect_columns(grid: list) -> tuple:
    """回傳 (描述欄, 原因欄, 處置欄, 資料起始列)。認不出來回 None。

    先找精確表頭組合，再用前五列關鍵字，找不到再退回「掃前 20 列，哪一欄最常以代碼開頭」。
    後者是為了沒有表頭或表頭被合併儲存格吃掉的情況。

    兩個實測踩到的坑：
      * FILL203 的 alarm list 分頁，Cause 欄完全沒有表頭（儲存格是 None），
        只靠關鍵字比對會整欄漏掉 → 表頭找不到 cause 時，退回用 desc+1。
      * Problem 分頁是問答清單不是警報資料，但內文提到代碼，會被 fallback
        誤判 → 要求至少 MIN_CODE_ROWS 列以代碼開頭才採用該分頁。
    """
    precise_row = _find_precise_header_row(grid)
    if precise_row is not None:
        cells = [_normalize_header_cell(value) for value in grid[precise_row]]
        # desc_i = Alarm Number 是階段一的權宜值，遷就 grid_to_rows() 既有的
        # 三角色（代碼欄/原因欄/處置欄）介面——Alarm Number 欄在這批檔案裡是
        # 純數字，split_code() 會把它整段當代碼、variant 永遠是空字串，
        # HMI Message 欄（真正的警報名稱文字）目前完全沒被這個回傳值引用。
        # 這不是刻意的欄位對應決策，是階段二完整欄位映射改寫前的過渡狀態
        # （見 line3_import_format_support_proposal.md 81、98、194 行），
        # 不要把這個回傳形狀當成定案抄去別處用。
        return (cells.index("alarm number"),
                cells.index("alarm description") if "alarm description" in cells else None,
                cells.index("solution") if "solution" in cells else None,
                precise_row + 2)

    HEAD = {"desc": ("description", "message", "alarm", "代碼", "訊息", "描述"),
            "cause": ("cause", "reason", "原因"),
            "action": ("action", "solution", "remedy", "comment", "處置", "對策", "備註")}
    MIN_CODE_ROWS = 5

    def code_rows(desc_i: int, start: int) -> int:
        return sum(1 for r in grid[start:]
                   if desc_i < len(r) and r[desc_i] and CODE_PREFIX_RE.match(str(r[desc_i])))

    for i, row in enumerate(grid[:5]):
        cells = [str(c).strip().lower() if c else "" for c in row]
        hit = {}
        for key, words in HEAD.items():
            for j, c in enumerate(cells):
                if c and any(w in c for w in words):
                    hit.setdefault(key, j)
        if "desc" not in hit:
            continue
        desc_i, start = hit["desc"], i + 1
        if code_rows(desc_i, start) < MIN_CODE_ROWS:
            continue
        cause_i, action_i = hit.get("cause"), hit.get("action")
        if cause_i is None:
            cand = desc_i + 1
            if cand != action_i and any(cand < len(r) and clean(r[cand]) for r in grid[start:]):
                cause_i = cand
        return desc_i, cause_i, action_i, start

    counts: Counter = Counter()
    for row in grid[:20]:
        for j, c in enumerate(row):
            if c and CODE_PREFIX_RE.match(str(c)):
                counts[j] += 1
    if not counts:
        return None
    desc_i = counts.most_common(1)[0][0]
    if code_rows(desc_i, 0) < MIN_CODE_ROWS:
        return None  # 零星提到代碼的分頁（例如問答清單），不是警報資料
    return desc_i, desc_i + 1, desc_i + 2, 0


def grid_to_rows(grid: list, mapping: tuple, source: str) -> list:
    """依 detect_columns() 的結果（或人工確認/修改過的對應）把 grid
    轉成 rows。code 切分（split_code()）在這裡才發生——這一步之前
    的所有資料都還是原始值，未被系統擅自解讀。"""
    desc_i, cause_i, action_i, start = mapping
    rows = []
    for r in grid[start:]:
        if desc_i >= len(r):
            continue
        code, variant = split_code(str(r[desc_i] or ""))
        if code is None:
            continue
        rows.append({
            "code": code,
            "variant": variant,
            "cause": clean(r[cause_i]) if cause_i is not None and cause_i < len(r) else "",
            "action": clean(r[action_i]) if action_i is not None and action_i < len(r) else "",
            "_source": source,
        })
    return rows


def grid_to_alarm_dicts(grid: list, header_row_idx: int, merges: list,
                        source: str, device_model: str, *, existing_keys: set) -> list:
    """依已確認的 LINE3 表頭轉出標準 alarm dict；不讀寫資料庫。

    header_row_idx 為 0-based，資料從表頭後第二列開始；merges 應由
    read_grid(..., with_merges=True) 提供。device_model 原樣保留。
    existing_keys 必填：第一批傳 set()，後續各檔／分頁傳入先前輸出的
    (device_model, code, variant) 集合；本函式不修改該集合。批內或跨批
    重複、非空但非純數字的 Alarm Number 都拋 ValueError，中止整批。

    每筆另含 _source 與 _warnings（沿用 Operating Modes 的警告結構，
    row/column 為 0-based，原始文字在 raw_value）。未知 Alarm Type
    的 severity 為 None 並附警告，呼叫端須呈現警告、處理後才能匯入。
    operating_modes 缺欄為 None，有欄但全未啟用為 []。
    """
    if not isinstance(device_model, str):
        raise TypeError("device_model 必須是字串")
    if not 0 <= header_row_idx < len(grid):
        raise ValueError("LINE3 表頭列索引超出範圍")
    header = grid[header_row_idx]
    normalized = [_normalize_header_cell(value) for value in header]
    columns = {}
    for field in ("alarm type", "alarm number", "hmi message",
                  "alarm description", "solution"):
        if normalized.count(field) != 1:
            raise ValueError(f"LINE3 表頭必須恰有一個 {field!r}")
        columns[field] = normalized.index(field)
    col_range = _find_operating_modes_range(header_row_idx, header, merges)
    sub_fields = (_extract_operating_modes_sub_fields(grid, header_row_idx, col_range)
                  if col_range is not None else [])
    seen = set(existing_keys)
    rows = []
    for row_idx in range(header_row_idx + 2, len(grid)):
        row = grid[row_idx]

        def cell(field):
            col = columns[field]
            return row[col] if col < len(row) else None

        code = _cell_to_str(cell("alarm number")).strip()
        if not code:
            continue
        if re.fullmatch(r"\d+", code) is None:
            raise ValueError(f"{source} row={row_idx}: 無效 Alarm Number {code!r}")
        key = (device_model, code, "")
        if key in seen:
            raise ValueError(f"{source} row={row_idx}: 重複 alarm 主鍵 {key!r}")
        seen.add(key)
        modes, warnings = extract_operating_modes(
            grid, header_row_idx, col_range, sub_fields, row_idx)
        alarm_type = cell("alarm type")
        severity = {"a": "警告", "w": "資訊"}.get(_normalize_header_cell(alarm_type))
        if severity is None:
            warnings.append({"field": "Alarm Type", "row": row_idx,
                             "column": columns["alarm type"],
                             "raw_value": _cell_to_str(alarm_type)})
        rows.append({
            "code": code, "variant": "", "device_model": device_model,
            "severity": severity, "description": clean(cell("hmi message")),
            "cause": clean(cell("alarm description")),
            "solution": clean(cell("solution")), "operating_modes": modes,
            "_source": source, "_warnings": warnings,
        })
    return rows


def read_tabular(path: pathlib.Path, sheet: str = None) -> list:
    """讀 .xlsx/.xlsm/.csv，自動掃描分頁（或 sheet 指定的單一分頁）並
    套用 detect_columns() 的建議。CLI 用這支——等同「人工確認欄位對應」
    這一步在 CLI 上是自動通過（CLI 使用者本來就知道自己在做什麼）。
    後台 inspect 端點不走這支，而是分開呼叫 read_grid()/detect_columns()，
    把建議交給人確認後才呼叫 grid_to_rows()。"""
    rows: list = []
    for name, grid, _ in read_grid(path, sheet=sheet):
        cols = detect_columns(grid)
        if cols is None:
            print(f"  [略過] 分頁 {name!r}：找不到含警報代碼的欄位", file=sys.stderr)
            continue
        n_before = len(rows)
        rows.extend(grid_to_rows(grid, cols, source=f"{path.name}#{name}"))
        print(f"  分頁 {name!r}：{len(rows) - n_before} 列", file=sys.stderr)
    return rows


def list_sheets(path: pathlib.Path) -> list:
    """回傳 [(分頁名, 是否偵測到警報欄位), ...]，供 --sheet-list 與
    inspect 端點列出分頁供人選擇，不解析內容。"""
    return [(name, detect_columns(grid) is not None) for name, grid, _ in read_grid(path)]
